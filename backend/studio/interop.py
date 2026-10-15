"""Getting decks in and out: embed the source in exported files, import decks, PDF conversion."""
from __future__ import annotations

import base64
import io
import json
import re
import shutil
import subprocess
import tempfile
import zlib
from pathlib import Path

import yaml
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

from . import templates as tpl
from .models import Meta, Proposal, Slide

CUSTOM_XML_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
NS = "urn:proposal-studio:source:1"
PART_NAME = "/customXml/proposal-studio1.xml"


class ImportError_(ValueError):
    pass


# ------------------------------------------------------------ round trip
def embed(prs, proposal: Proposal) -> None:
    """Store the proposal inside the .pptx as a standard customXml item (ignored by PowerPoint)."""
    payload = base64.b64encode(zlib.compress(proposal.model_dump_json().encode("utf-8"))).decode("ascii")
    xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<proposalStudio xmlns="{NS}" version="1">{payload}</proposalStudio>'
    for rel in list(prs.part.rels.values()):  # replace an item left by a previous export
        if rel.reltype == CUSTOM_XML_RT and not rel.is_external and str(rel.target_part.partname) == PART_NAME:
            prs.part.drop_rel(rel.rId)
    part = Part(PackURI(PART_NAME), "application/xml", prs.part.package, xml.encode("utf-8"))
    prs.part.relate_to(part, CUSTOM_XML_RT)


def extract(prs) -> Proposal | None:
    for rel in prs.part.rels.values():
        if rel.reltype != CUSTOM_XML_RT or rel.is_external:
            continue
        blob = rel.target_part.blob
        if NS.encode() not in blob[:400]:
            continue
        try:
            root = etree.fromstring(blob)
            raw = zlib.decompress(base64.b64decode(root.text or ""))
            return Proposal.model_validate_json(raw)
        except Exception:
            return None
    return None


# ------------------------------------------------------- approximate import
def _lines(shape) -> list[str]:
    out: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sh in shape.shapes:
            out += _lines(sh)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(dict.fromkeys(cells)))
    elif shape.has_text_frame:
        out += [p.strip() for p in shape.text_frame.text.splitlines() if p.strip()]
    return out


def from_pptx(data: bytes) -> tuple[Proposal, list[str], str]:
    """Return (proposal, warnings, source). Exact when the file was exported by this app."""
    try:
        clean = tpl.normalize(data)
        prs = Presentation(io.BytesIO(clean))
    except tpl.TemplateError as e:
        raise ImportError_(str(e)) from e
    except Exception as e:  # python-pptx raises many types on corrupt files
        raise ImportError_(f"could not read the file: {type(e).__name__}") from e
    exact = extract(prs)
    if exact is not None:
        return exact, [], "embedded"

    slides: list[Slide] = []
    for idx, sl in enumerate(prs.slides):
        title = sl.shapes.title.text_frame.text.strip() if sl.shapes.title is not None and sl.shapes.title.has_text_frame else ""
        body: list[str] = []
        for sh in sl.shapes:
            if sh is sl.shapes.title:
                continue
            if sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE):
                continue
            body += [t for t in _lines(sh) if not re.fullmatch(r"\d{1,3}", t)]
        if not title and body:
            title, body = body[0][:120], body[1:]
        notes = ""
        if sl.has_notes_slide:
            notes = sl.notes_slide.notes_text_frame.text.strip()
        if idx == 0 and len(body) <= 2:
            slides.append(Slide(type="title", title=title, bullets=body, notes=notes))
        else:
            slides.append(Slide(type="bullets", title=title or f"Slide {idx + 1}", bullets=body[:25], notes=notes))
    if not slides:
        raise ImportError_("the presentation has no slides")
    first = slides[0]
    meta = Meta(title=(first.title or prs.core_properties.title or "Imported deck")[:200],
                subtitle=(first.bullets[0] if first.type == "title" and first.bullets else "")[:200])
    slides = [Slide(type="title", notes=first.notes)] + slides[1:] if first.type == "title" else slides
    warnings = ["This file was not exported by Proposal Studio, so its text was imported as editable bullet slides. "
                "Diagrams, charts and pictures are not converted."]
    return Proposal(meta=meta, slides=slides), warnings, "approximate"


def from_text(data: bytes) -> tuple[Proposal, list[str], str]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise ImportError_("the file is not UTF-8 text, .pptx, .json or .yaml") from e
    try:
        obj = json.loads(text)
        src = "json"
    except json.JSONDecodeError:
        try:
            obj = yaml.safe_load(text)
            src = "yaml"
        except yaml.YAMLError as e:
            raise ImportError_("could not parse the file as JSON or YAML") from e
    if not isinstance(obj, dict):
        raise ImportError_("expected a proposal object with `meta` and `slides`")
    from pydantic import ValidationError
    try:
        return Proposal.model_validate(obj), [], src
    except ValidationError as e:
        first = e.errors()[0]
        raise ImportError_(f"not a valid proposal: {'.'.join(str(x) for x in first['loc'])}: {first['msg']}") from e


def _looks_like_html(data: bytes) -> bool:
    head = data[:4096].lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    return head.startswith((b"<!doctype html", b"<html", b"<head", b"<body", b"<section", b"<div", b"<!--", b"<article", b"<h1"))


def import_any(data: bytes) -> tuple[Proposal, list[str], str]:
    if data[:2] == b"PK":
        return from_pptx(data)
    if _looks_like_html(data):
        from . import htmlimport
        try:
            return htmlimport.from_html(data)
        except ValueError as e:
            raise ImportError_(str(e)) from e
    return from_text(data)


# ------------------------------------------------------------------- pdf
def soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def to_pdf(pptx_path: Path, timeout: int = 120) -> Path:
    exe = soffice_path()
    if not exe:
        raise RuntimeError("LibreOffice is not installed on the server, so PDF export is unavailable")
    out_dir = Path(tempfile.mkdtemp(prefix="studio-pdf-"))
    profile = Path(tempfile.mkdtemp(prefix="studio-lo-"))
    try:
        subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(out_dir), str(pptx_path)], check=True, timeout=timeout,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except subprocess.TimeoutExpired as e:
        raise RuntimeError("PDF conversion timed out") from e
    except subprocess.CalledProcessError as e:
        raise RuntimeError("PDF conversion failed") from e
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    pdf = out_dir / (pptx_path.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("PDF conversion produced no file")
    return pdf
