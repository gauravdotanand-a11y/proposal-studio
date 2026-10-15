"""Render a Proposal to an editable .pptx using native shapes, tables and connectors."""
from __future__ import annotations

import copy
import os
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import templates as TPL
from . import theme as T
from .fit import expand, gantt_row_heights, gantt_rows, table_row_h, title_size, est_lines, text_h
from . import interop
from .layout import layout_flow
from .models import Assumptions, Cost, Flow, Gantt, Proposal, Slide

_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _text(tf, text, size=None, bold=False, color=None, align=None):
    size = T.BODY_SIZE if size is None else size
    color = T.GREY if color is None else color  # resolved per call so looks and templates can override
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.font.size, r.font.bold = size, bold
        if T.FONT:
            r.font.name = T.FONT
        r.font.color.rgb = color
    return p


def _set_ph_text(shape, text):
    shape.text_frame.text = text


def _title(slide, text):
    if slide.shapes.title is not None:
        _set_ph_text(slide.shapes.title, text)
        if len(text) > 55:  # placeholder size comes from the layout; step it down for long titles
            for p in slide.shapes.title.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(24 if len(text) <= 85 else 20)
        return
    width = T.SLIDE_W - 2 * T.MARGIN
    size = title_size(text, width, base=T.TITLE_SIZE / 12700)
    tb = slide.shapes.add_textbox(T.MARGIN, Inches(0.35), width, Inches(0.8))
    _text(tb.text_frame, text, Pt(size), True, T.NAVY)


def _template_footer(slide, layout, meta, number) -> bool:
    """Clone the layout's footer / slide-number placeholders so the master's branding shows."""
    found = False
    for ph in layout.placeholders:
        kind = ph.placeholder_format.type
        if kind not in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
            continue
        el = copy.deepcopy(ph._element)
        el.nvSpPr.cNvPr.set("id", str(slide.shapes._next_shape_id))
        slide.shapes._spTree.append(el)
        shape = slide.shapes[-1]
        has_field = "slidenum" in etree.tostring(el).decode()
        if kind == PP_PLACEHOLDER.FOOTER:
            shape.text_frame.text = meta.footer or meta.brand
        elif not has_field:
            shape.text_frame.text = str(number)
        found = True
    return found


def _footer(slide, meta, number):
    tb = slide.shapes.add_textbox(T.MARGIN, T.SLIDE_H - Inches(0.5), Inches(3), Inches(0.35))
    _text(tb.text_frame, meta.brand, Pt(14), True, T.ORANGE)
    tb = slide.shapes.add_textbox(Inches(4), T.SLIDE_H - Inches(0.45), Inches(5.3), Inches(0.3))
    _text(tb.text_frame, meta.footer, Pt(9), False, T.GREY, PP_ALIGN.CENTER)
    tb = slide.shapes.add_textbox(T.SLIDE_W - Inches(1.2), T.SLIDE_H - Inches(0.45), Inches(0.6), Inches(0.3))
    _text(tb.text_frame, str(number), Pt(10), False, T.GREY, PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- slides
def _title_slide(prs, slide, meta):
    if slide.shapes.title is not None:
        _set_ph_text(slide.shapes.title, meta.title)
        for ph in slide.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                _set_ph_text(ph, meta.subtitle)
        return
    width = T.SLIDE_W - 2 * T.MARGIN
    size = next((z for z in (40, 34, 30, 26, 22) if est_lines(meta.title, width, z, True) <= 3), 20)
    h = text_h(est_lines(meta.title, width, size, True), size) + Inches(0.1)
    top = Inches(2.6) if h <= Inches(1.3) else max(Inches(1.4), Inches(3.2) - h / 2)
    tb = slide.shapes.add_textbox(T.MARGIN, int(top), width, int(h))
    _text(tb.text_frame, meta.title, Pt(size), True, T.NAVY)
    if meta.subtitle:
        sub = next((z for z in (20, 18, 16, 14) if est_lines(meta.subtitle, width, z) <= 3), 14)
        tb = slide.shapes.add_textbox(T.MARGIN, int(top + h + Inches(0.15)), width, Inches(0.8))
        _text(tb.text_frame, meta.subtitle, Pt(sub), False, T.GREY)


def _bullets_slide(slide, title, items, numbered=False, size=15, start=0):
    _title(slide, title)
    tb = slide.shapes.add_textbox(T.MARGIN, T.CONTENT_TOP, T.SLIDE_W - 2 * T.MARGIN, T.SLIDE_H - T.CONTENT_TOP - Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{start + i + 1}. {item}" if numbered else f"• {item}"
        p.space_after = Pt(8)
        for r in p.runs:
            r.font.size, r.font.color.rgb = Pt(size), T.GREY
            if T.FONT:
                r.font.name = T.FONT


def _cell(cell, text, fill=None, bold=False, color=None, size=Pt(10), align=PP_ALIGN.CENTER):
    color = T.GREY if color is None else color
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    for r in p.runs:
        r.font.size, r.font.bold = size, bold
        if T.FONT:
            r.font.name = T.FONT
        r.font.color.rgb = color
    p._p.get_or_add_endParaRPr().set("sz", str(int(round(size / 12700 * 100))))
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.04)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def _gantt_table(slide, g: Gantt, left, top, width, row_h=None, bar=None, changed=frozenset(), compact=False):
    """Week grid as an editable table. Section headers become full-width bands.

    Returns (graphic frame, total height in EMU); row heights follow the wrapped activity names.
    """
    n = g.week_count
    name_w = Inches(2.9 if compact else 3.4)
    rows = [("head1", None), ("head2", None), *gantt_rows(g.activities)]
    row_h = int(row_h or Inches(0.4))
    body_h = gantt_row_heights(g.activities, row_h, compact)
    heights = [row_h, row_h] + body_h
    gf = slide.shapes.add_table(len(rows), n + 1, left, top, width, sum(heights))
    tbl = gf.table
    tbl.columns[0].width = name_w
    wk_w = int((width - name_w) / n)
    for c in range(1, n + 1):
        tbl.columns[c].width = wk_w
    for r, h in enumerate(heights):
        tbl.rows[r].height = int(h)
    f_head, f_dates, f_name = (Pt(9), Pt(6), Pt(9)) if compact else (Pt(10), Pt(7), Pt(10))
    dense = wk_w < Inches(0.55)  # many weeks: short labels so headers do not wrap
    if dense:
        f_head, f_dates = Pt(7), Pt(5)
    bar = bar or T.ORANGE
    for r, (kind, payload) in enumerate(rows):
        if kind == "head1":
            _cell(tbl.cell(r, 0), "", T.NAVY)
            for w in range(1, n + 1):
                _cell(tbl.cell(r, w), f"W{w}" if dense else f"week {w}", T.NAVY, True, T.WHITE, f_head)
        elif kind == "head2":
            _cell(tbl.cell(r, 0), "Activities", T.NAVY, True, T.WHITE, f_head + Pt(1))
            for w in range(1, n + 1):
                s_, e_ = g.week_range(w)
                label = f"{s_.day} {s_:%b}" if dense else f"{s_.day} {s_:%b} - {e_.day} {e_:%b}"
                _cell(tbl.cell(r, w), label, T.TEAL, False, T.NAVY, f_dates)
        elif kind == "section":
            tbl.cell(r, 0).merge(tbl.cell(r, n))
            _cell(tbl.cell(r, 0), str(payload), T.TEAL_DARK, True, T.NAVY, f_name, PP_ALIGN.LEFT)
        else:
            a = payload
            fill = T.NOTE_FILL if a.name in changed else T.TEAL
            _cell(tbl.cell(r, 0), a.name, fill, a.name in changed, T.NAVY, f_name, PP_ALIGN.LEFT)
            for w in range(1, n + 1):
                _cell(tbl.cell(r, w), "", bar if a.start_week <= w <= a.end_week else None)
    return gf, sum(heights)


def _gantt_slide(slide, g: Gantt, row_h=None):
    _title(slide, g.title)
    _gantt_table(slide, g, T.MARGIN, T.CONTENT_TOP + Inches(0.1), T.SLIDE_W - 2 * T.MARGIN, row_h=row_h or Inches(0.4))
    if g.note:
        tb = slide.shapes.add_textbox(T.MARGIN, T.SLIDE_H - Inches(1.3), T.SLIDE_W - 2 * T.MARGIN, Inches(0.4))
        _text(tb.text_frame, g.note, Pt(11), True, T.NAVY)


def _compare_slide(slide, c):
    _title(slide, c.title)
    width = T.SLIDE_W - 2 * T.MARGIN
    y = T.CONTENT_TOP - Inches(0.1)
    footer_h = Inches(0.55)
    avail = T.SLIDE_H - Inches(0.8) - y - footer_h - 2 * Inches(0.4)
    total_rows = sum(2 + len(gantt_rows(g.activities)) for g in (c.previous, c.updated)) - 4
    row_h = max(Inches(0.22), min(Inches(0.3), int(avail / max(1, total_rows + 4))))
    changed = frozenset(c.changed)
    for label, g, bar, hi in ((c.previous_label, c.previous, T.TEAL_DARK, frozenset()),
                              (c.updated_label, c.updated, T.ORANGE, changed)):
        tb = slide.shapes.add_textbox(T.MARGIN, y, width, Inches(0.35))
        _text(tb.text_frame, label, Pt(15), True, T.NAVY)
        y += Inches(0.38)
        gf, h = _gantt_table(slide, g, T.MARGIN, y, width, row_h=row_h, bar=bar, changed=hi, compact=True)
        gf.name = f"gantt:{label}"
        y += h + Inches(0.1)
    d = c.delta_weeks
    summary = (f"Duration {'+' if d > 0 else ''}{d} week(s) vs previous; " if d else "Same duration as previous; ")
    summary += f"{len(changed)} activit{'y' if len(changed) == 1 else 'ies'} changed (highlighted)."
    tb = slide.shapes.add_textbox(T.MARGIN, min(y + Inches(0.05), T.SLIDE_H - Inches(0.95)), width, Inches(0.4))
    _text(tb.text_frame, summary, Pt(11), True, T.NAVY)


# --------------------------------------------------------- scope / lists
def _scope_slide(slide, sc, size=13):
    _title(slide, sc.title)
    col_gap = Inches(0.4)
    width = T.SLIDE_W - 2 * T.MARGIN
    col_w = int((width - col_gap) / 2)
    top = T.CONTENT_TOP
    for i, (head, items, fill, ink) in enumerate((
            ("In scope", sc.in_scope, T.TEAL, T.NAVY), ("Out of scope", sc.out_scope, T.OUT_FILL, T.GREY))):
        x = int(T.MARGIN + i * (col_w + col_gap))
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top, col_w, Inches(0.45))
        band.name = f"scope-head:{i}"
        band.fill.solid()
        band.fill.fore_color.rgb = fill
        band.line.fill.background()
        band.shadow.inherit = False
        band.text_frame.margin_left = Inches(0.15)
        band.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(band.text_frame, head, Pt(15), True, ink, PP_ALIGN.LEFT)
        body = slide.shapes.add_textbox(x, top + Inches(0.55), col_w, T.SLIDE_H - top - Inches(1.9))
        tf = body.text_frame
        tf.word_wrap = True
        for k, it in enumerate(items):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = f"\u2022 {it}"
            p.space_after = Pt(6)
            for r in p.runs:
                r.font.size, r.font.color.rgb = Pt(size), T.GREY
                if T.FONT:
                    r.font.name = T.FONT
    if sc.note:
        tb = slide.shapes.add_textbox(T.MARGIN, T.SLIDE_H - Inches(1.3), width, Inches(0.4))
        _text(tb.text_frame, sc.note, Pt(11), True, T.NAVY)


def _simple_table(slide, headers, rows, widths, fills=None):
    width = T.SLIDE_W - 2 * T.MARGIN
    total = sum(widths)
    cols = [width * w / total for w in widths]
    heights = [Inches(0.45)] + [table_row_h([(str(t), cols[c], 11, False) for c, t in enumerate(r)], Inches(0.32)) for r in rows]
    gf = slide.shapes.add_table(len(rows) + 1, len(headers), T.MARGIN, T.CONTENT_TOP, width, sum(heights))
    tbl = gf.table
    for i, w in enumerate(cols):
        tbl.columns[i].width = int(w)
    for r, h in enumerate(heights):
        tbl.rows[r].height = int(h)
    for i, h in enumerate(headers):
        _cell(tbl.cell(0, i), h, T.NAVY, True, T.WHITE, Pt(12), PP_ALIGN.LEFT)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            fill, ink, bold = (fills or {}).get((r - 1, c), (None, T.GREY, False))
            _cell(tbl.cell(r, c), val, fill, bold, ink, Pt(11), PP_ALIGN.LEFT)
    return gf


def _risks_slide(slide, rk):
    _title(slide, rk.title)
    fills = {}
    for r, it in enumerate(rk.items):
        fills[(r, 1)] = {"High": (T.ORANGE, T.WHITE, True), "Medium": (T.NOTE_FILL, T.NAVY, True),
                         "Low": (T.TEAL, T.NAVY, True)}[it.impact]
    gf = _simple_table(slide, ["Risk / dependency", "Impact", "Mitigation"],
                       [[i.risk, i.impact, i.mitigation] for i in rk.items], [5, 1.4, 5], fills)
    gf.name = "risks-table"


def _next_steps_slide(slide, ns, start=0):
    _title(slide, ns.title)
    gf = _simple_table(slide, ["#", "Action", "Owner", "Due"],
                       [[str(i), s.action, s.owner, s.due] for i, s in enumerate(ns.items, start=start + 1)], [0.5, 7, 2.2, 1.8])
    gf.name = "next-steps-table"


def _money(cur, v):
    return f"{cur}{v:,.2f}".rstrip("0").rstrip(".") if v else "-"


def _cost_slide(slide, c: Cost, row_h=None, lines=None, final=True):
    """Cost table, then (on the final part) CBA, benefits and driver basis.

    `lines` selects which cost lines this part shows; totals always cover every line.
    """
    _title(slide, c.title)
    cols = len(c.years) + 2
    width = T.SLIDE_W - 2 * T.MARGIN
    row_h = int(row_h or Inches(0.3))
    show = [c.lines[i] for i in lines] if lines is not None else c.lines
    name_w = Inches(3.2)
    heights = [Inches(0.45)] + [table_row_h([(ln.head, name_w, 11, False)], row_h) for ln in show]
    if final:
        heights.append(Inches(0.45))
    rows = len(heights)
    gf = slide.shapes.add_table(rows, cols, T.MARGIN, T.CONTENT_TOP, width, sum(heights))
    gf.name = "cost-table"
    tbl = gf.table
    tbl.columns[0].width = name_w
    other = int((width - name_w) / (cols - 1))
    for i in range(1, cols):
        tbl.columns[i].width = other
    for r, h in enumerate(heights):
        tbl.rows[r].height = int(h)
    heads = ["Cost Head", *c.years, "Total"]
    for i, h in enumerate(heads):
        _cell(tbl.cell(0, i), h, T.NAVY, True, T.WHITE, Pt(12))
    for r, line in enumerate(show, start=1):
        _cell(tbl.cell(r, 0), line.head, None, False, T.GREY, Pt(11), PP_ALIGN.LEFT)
        for i, v in enumerate(line.values, start=1):
            _cell(tbl.cell(r, i), _money(c.currency, v), None, False, T.GREY, Pt(11))
        _cell(tbl.cell(r, cols - 1), _money(c.currency, sum(line.values)), None, True, T.GREY, Pt(11))
    if not final:
        return
    last = rows - 1
    _cell(tbl.cell(last, 0), "Total", T.ORANGE, True, T.WHITE, Pt(12), PP_ALIGN.LEFT)
    for i, v in enumerate(c.year_totals, start=1):
        _cell(tbl.cell(last, i), _money(c.currency, v), T.ORANGE, True, T.WHITE, Pt(12))
    _cell(tbl.cell(last, cols - 1), _money(c.currency, c.total_cost), T.ORANGE, True, T.WHITE, Pt(12))

    y = T.CONTENT_TOP + sum(heights) + Inches(0.3)
    badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, T.MARGIN, y, Inches(4), Inches(0.5))
    badge.fill.background()
    badge.line.color.rgb = T.NAVY
    badge.text_frame.text = ""
    _text(badge.text_frame, f"CBA - {c.cba:.1f}x", Pt(18), True, T.NAVY, PP_ALIGN.CENTER)

    if c.one_time_total or c.recurring_total:
        info = slide.shapes.add_textbox(T.MARGIN + Inches(4.2), y, width - Inches(4.2), Inches(0.5))
        _text(info.text_frame, f"One-time: {_money(c.currency, c.one_time_total)}   |   "
                               f"Recurring (all years): {_money(c.currency, c.recurring_total)}",
              Pt(11), True, T.NAVY, PP_ALIGN.LEFT)
    brow = slide.shapes.add_table(2, cols, T.MARGIN, y + Inches(0.8), width, Inches(0.8)).table
    brow.columns[0].width = name_w
    for i in range(1, cols):
        brow.columns[i].width = other
    for i, h in enumerate(["Benefits", *c.years, "Total"]):
        _cell(brow.cell(0, i), h, T.NAVY, True, T.WHITE, Pt(12), PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER)
    _cell(brow.cell(1, 0), "", None)
    for i, v in enumerate(c.benefit_values, start=1):
        _cell(brow.cell(1, i), _money(c.currency, v), None, False, T.GREY, Pt(11))
    _cell(brow.cell(1, cols - 1), _money(c.currency, c.total_benefit), None, True, T.GREY, Pt(11))
    if c.drivers:
        basis = slide.shapes.add_textbox(T.MARGIN, y + Inches(1.75), width, Inches(0.9))
        tf = basis.text_frame
        tf.word_wrap = True
        for k, d in enumerate(c.drivers):
            line = (f"{d.name}: {d.annual_volume:,.0f} items/yr x {d.minutes_saved:g} min x "
                    f"{c.currency}{d.hourly_rate:g}/h x {d.success_rate:.0%} success")
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = line
            for r in p.runs:
                r.font.size, r.font.color.rgb = Pt(9), T.GREY
                if T.FONT:
                    r.font.name = T.FONT


# ------------------------------------------------------------------ flow
_SHAPE = {
    "decision": MSO_SHAPE.DIAMOND,
    "start": MSO_SHAPE.OVAL,
    "end": MSO_SHAPE.OVAL,
    "data": MSO_SHAPE.CAN,
}


def _set_descr(shape, text: str) -> None:
    """Alt text for screen readers (the shape's cNvPr `descr`)."""
    for el in shape._element.iter(f"{{{_NS_P}}}cNvPr"):
        el.set("descr", text[:250])
        break


def _describe_tables(slide, s: Slide) -> None:
    what = {"gantt": "Gantt chart of activities by week", "compare": "Gantt chart comparing plan versions",
            "cost": "Cost and benefit table", "risks": "Table of risks, impact and mitigation",
            "next_steps": "Table of next steps with owner and due date"}.get(s.type)
    if not what:
        return
    for sh in slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table and not sh._element.xpath(".//p:cNvPr[@descr]"):
            _set_descr(sh, f"{what}: {sh.table.cell(0, 0).text or ''} {len(sh.table.rows) - 1} rows".strip())


def _flow_slide(slide, f: Flow):
    _title(slide, f.title)
    lay = layout_flow(f, top=T.CONTENT_TOP + Inches(0.1))
    for i, ln in enumerate(lay.get("lanes", [])):
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ln["x"], ln["y"], ln["w"], ln["h"])
        band.name = f"lane:{ln['name']}"
        band.fill.solid()
        band.fill.fore_color.rgb = T.WHITE if i % 2 == 0 else T.LANE_TINT
        band.line.color.rgb = T.TEAL_DARK
        band.line.width = Pt(0.75)
        band.shadow.inherit = False
        lab = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ln["x"], ln["y"], ln["label_w"], ln["h"])
        lab.name = f"lane-label:{ln['name']}"
        lab.fill.solid()
        lab.fill.fore_color.rgb = T.TEAL
        lab.line.color.rgb = T.TEAL_DARK
        lab.line.width = Pt(0.75)
        lab.shadow.inherit = False
        lab.text_frame.margin_left = lab.text_frame.margin_right = Inches(0.05)
        lab.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(lab.text_frame, ln["name"], Pt(11), True, T.NAVY, PP_ALIGN.CENTER)
    shapes = {}
    badges = []
    for n in lay["nodes"]:
        shp = slide.shapes.add_shape(_SHAPE.get(n["kind"], MSO_SHAPE.ROUNDED_RECTANGLE),
                                     n["x"], n["y"], n["w"], n["h"])
        fill = {"bot": T.TEAL_DARK, "ai": T.ORANGE, "human": T.WHITE}.get(n["kind"], T.TEAL)
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.line.color.rgb = T.NAVY if n["kind"] in ("human", "decision") else fill
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        size = Pt(9) if n["w"] < Inches(1.4) or n["kind"] == "decision" else Pt(10)
        if n["kind"] == "decision":
            tf.margin_left = tf.margin_right = 0
        _text(tf, n["label"], size, False, T.WHITE if n["kind"] == "ai" else T.NAVY, PP_ALIGN.CENTER)
        shp.name = f"node:{n['id']}"
        _set_descr(shp, f"{n['kind']} step: {n['label']}")
        shapes[n["id"]] = shp
        if n.get("badge"):
            badges.append(n)

    kinds = {n["id"]: n["kind"] for n in lay["nodes"]}
    # PowerPoint connection-site index differs per shape: ellipses have 8 sites, cylinders 5.
    site_idx = {"start": {0: 0, 1: 2, 2: 4, 3: 6}, "end": {0: 0, 1: 2, 2: 4, 3: 6},
                "data": {0: 1, 1: 2, 2: 3, 3: 4}}
    for e in lay["edges"]:
        pts = e["points"]
        glued = e["begin_site"] is not None and e["end_site"] is not None
        elbow = (len(pts) == 4 and glued and pts[0][1] == pts[1][1] and pts[1][0] == pts[2][0]
                 and pts[2][1] == pts[3][1] and pts[3][0] != pts[0][0])
        corner = (len(pts) == 3 and glued and pts[0][1] == pts[1][1] and pts[1][0] == pts[2][0]
                  and pts[2][0] != pts[0][0] and pts[2][1] != pts[1][1])  # horizontal leg, then vertical
        if glued and (len(pts) == 2 or elbow or corner):
            kind = MSO_CONNECTOR.STRAIGHT if len(pts) == 2 else MSO_CONNECTOR.ELBOW
            cx = slide.shapes.add_connector(kind, 0, 0, 0, 0)
            cx.begin_connect(shapes[e["src"]], e["begin_site"])
            cx.end_connect(shapes[e["dst"]], e["end_site"])
            for tag, node, site in (("stCxn", e["src"], e["begin_site"]), ("endCxn", e["dst"], e["end_site"])):
                remap = site_idx.get(kinds[node])
                if remap:
                    el = cx._element.nvCxnSpPr.cNvCxnSpPr.find(f"{{{_NS_A}}}{tag}")
                    if el is not None:
                        el.set("idx", str(remap[site]))
            if corner:
                cx._element.spPr.find(f"{{{_NS_A}}}prstGeom").set("prst", "bentConnector2")
            if elbow:  # place the vertical leg where the layout routed it (fraction of the width)
                frac = (pts[1][0] - pts[0][0]) / (pts[3][0] - pts[0][0])
                geom = cx._element.spPr.find(f"{{{_NS_A}}}prstGeom")
                av = geom.find(f"{{{_NS_A}}}avLst")
                if av is None:
                    av = etree.SubElement(geom, f"{{{_NS_A}}}avLst")
                etree.SubElement(av, f"{{{_NS_A}}}gd", name="adj1", fmla=f"val {int(round(frac * 100000))}")
            line_owner = cx
        else:
            fb = slide.shapes.build_freeform(pts[0][0], pts[0][1])
            fb.add_line_segments([(x, y) for x, y in pts[1:]], close=False)
            line_owner = fb.convert_to_shape()
            line_owner.fill.background()
            line_owner.shadow.inherit = False
        line_owner.name = f"edge:{e['src']}->{e['dst']}"
        names = {n["id"]: n["label"] for n in lay["nodes"]}
        _set_descr(line_owner, f"Arrow from {names.get(e['src'], e['src'])} to {names.get(e['dst'], e['dst'])}"
                               f"{' (' + e['label'] + ')' if e['label'] else ''}")
        line_owner.line.color.rgb = T.NAVY
        line_owner.line.width = Pt(1.25)
        ln = line_owner.line._get_or_add_ln()
        etree.SubElement(ln, f"{{{_NS_A}}}tailEnd", type="triangle")
        if e["label"]:
            lb = slide.shapes.add_textbox(e["label_x"] - Inches(0.17), e["label_y"], Inches(0.34), Inches(0.25))
            lb.text_frame.margin_left = lb.text_frame.margin_right = 0
            lb.text_frame.word_wrap = False
            _text(lb.text_frame, e["label"], Pt(9), True, T.NAVY, PP_ALIGN.CENTER)

    for e in lay["edges"]:
        if e.get("note_box"):
            x, y, w, h = e["note_box"]
            tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
            tag.name = f"note:{e['src']}->{e['dst']}"
            tag.fill.solid()
            tag.fill.fore_color.rgb = T.NOTE_FILL
            tag.line.fill.background()
            tag.shadow.inherit = False
            tf = tag.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.03)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            _text(tf, e["note"], Pt(8), True, T.NAVY, PP_ALIGN.CENTER)
    for n in badges:  # drawn last so they sit above boxes and arrows
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, n["x"] + Inches(0.06), n["y"] - Inches(0.1), Inches(0.62), Inches(0.2))
        pill.name = f"badge:{n['id']}"
        pill.fill.solid()
        pill.fill.fore_color.rgb = T.ORANGE if n["kind"] == "ai" else T.NAVY
        pill.line.fill.background()
        pill.shadow.inherit = False
        ptf = pill.text_frame
        ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = 0
        ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(ptf, n["badge"], Pt(7), True, T.WHITE, PP_ALIGN.CENTER)

    ny = lay["notes_top"]
    for note in f.notes:
        tb = slide.shapes.add_textbox(T.MARGIN, ny, T.SLIDE_W - 2 * T.MARGIN, Inches(0.3))
        _text(tb.text_frame, "* " + note, Pt(10), False, T.GREY)
        ny += Inches(0.3)


# -------------------------------------------------------------- solution
def _rich(tf, parts, size, first=True):
    """Write one paragraph made of (text, bold) runs."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    for text, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold = size, bold
        r.font.color.rgb = T.NAVY if bold else T.GREY
        if T.FONT:
            r.font.name = T.FONT
    p.space_after = Pt(4)
    return p


def _solution_slide(slide, sol):
    _title(slide, sol.title)
    n = len(sol.stages)
    left, width = T.MARGIN, T.SLIDE_W - 2 * T.MARGIN
    arrow_gap = Inches(0.4)
    box_w = int((width - (n - 1) * arrow_gap) / n)
    head_h, box_h = Inches(0.35), Inches(1.55)
    top = T.CONTENT_TOP
    boxes = []
    for i, st in enumerate(sol.stages):
        x = int(left + i * (box_w + arrow_gap))
        if st.heading:
            hd = slide.shapes.add_textbox(x, top, box_w, head_h)
            _text(hd.text_frame, st.heading, Pt(12), True, T.NAVY, PP_ALIGN.CENTER)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top + head_h, box_w, box_h)
        box.name = f"stage:{i}"
        box.fill.solid()
        box.fill.fore_color.rgb = T.ORANGE if st.ai else T.TEAL
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ink = T.WHITE if st.ai else T.NAVY
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = st.title
        r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, ink
        if T.FONT:
            r.font.name = T.FONT
        if st.body:
            q = tf.add_paragraph()
            q.alignment = PP_ALIGN.CENTER
            r2 = q.add_run()
            r2.text = st.body
            r2.font.size, r2.font.color.rgb = Pt(9), ink
            if T.FONT:
                r2.font.name = T.FONT
        boxes.append(box)
    for a, b in zip(boxes, boxes[1:]):
        cx = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
        cx.begin_connect(a, 3)
        cx.end_connect(b, 1)
        cx.name = "stage-arrow"
        cx.line.color.rgb = T.NAVY
        cx.line.width = Pt(1.5)
        etree.SubElement(cx.line._get_or_add_ln(), f"{{{_NS_A}}}tailEnd", type="triangle")

    y0 = top + head_h + box_h + Inches(0.3)
    sep = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left, y0, left + width, y0)
    sep.line.color.rgb = T.TEAL_DARK
    sep.line.width = Pt(0.75)
    y0 += Inches(0.12)
    avail_h = T.SLIDE_H - Inches(0.8) - y0
    total = sum(len(b.text) + len(b.heading) for b in sol.blocks), sum(len(s.text) + len(s.label) for s in sol.steps)
    size = Pt(11) if max(total) < 600 else Pt(10) if max(total) < 950 else Pt(9)
    col_gap = Inches(0.4)
    left_w = int(width * 0.46)
    right_w = width - left_w - col_gap

    if sol.blocks:
        tb = slide.shapes.add_textbox(left, y0, left_w, avail_h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, blk in enumerate(sol.blocks):
            _rich(tf, [(blk.heading, True)], size + Pt(1), first=(i == 0))
            _rich(tf, [(blk.text, False)], size, first=False)
    if sol.steps:
        tb = slide.shapes.add_textbox(left + left_w + col_gap, y0, right_w, avail_h)
        tf = tb.text_frame
        tf.word_wrap = True
        _rich(tf, [(sol.steps_title, True)], size + Pt(1))
        for i, st in enumerate(sol.steps, start=1):
            parts = [(f"{i}. ", True)]
            if st.label:
                parts.append((st.label + ": ", True))
            parts.append((st.text, False))
            _rich(tf, parts, size, first=False)


# ------------------------------------------------------------------ main
def _strip_slides(prs) -> None:
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def _drop_empty_placeholders(slide) -> None:
    for ph in list(slide.placeholders):
        kind = ph.placeholder_format.type
        if kind in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            continue
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def render(proposal: Proposal, out: str | Path, template: str | Path | None = None,
           title_layout: int | None = None, content_layout: int | None = None,
           use_theme_colors: bool = True, warnings: list[str] | None = None,
           preset: str | None = None) -> Path:
    warnings = warnings if warnings is not None else []
    overrides: dict = {}
    if preset and not template:
        if preset not in T.PRESETS:
            raise ValueError(f"unknown look '{preset}'")
        overrides.update(T.preset_overrides(preset))
    if template:
        prs = Presentation(str(template))
        _strip_slides(prs)
        auto_title, auto_content = TPL.pick_layouts(prs)
        n_layouts = len(prs.slide_layouts)
        for label, idx in (("title", title_layout), ("content", content_layout)):
            if idx is not None and not 0 <= idx < n_layouts:
                raise ValueError(f"{label} layout index {idx} is out of range (0-{n_layouts - 1})")
        t_idx = auto_title if title_layout is None else title_layout
        c_idx = auto_content if content_layout is None else content_layout
        left, top = TPL.content_top_for(prs, c_idx)
        overrides.update(SLIDE_W=prs.slide_width, SLIDE_H=prs.slide_height, MARGIN=left, CONTENT_TOP=top, FONT=None)
        if use_theme_colors:
            pal = TPL.palette(TPL.theme_colors(prs))
            if pal:
                from pptx.dml.color import RGBColor
                overrides.update({k: RGBColor.from_string(v) for k, v in pal.items()})
            else:
                warnings.append("Template theme has no usable accent colours; built-in palette used.")
        w_in, h_in = prs.slide_width / 914400, prs.slide_height / 914400
        if abs(w_in / h_in - 16 / 9) > 0.03:
            warnings.append(f"Template is {w_in:.2f}x{h_in:.2f} in (not 16:9); diagrams and tables may be cramped.")
        if top > Inches(1.9):
            warnings.append("Template title area is tall; less room is left for diagrams and tables.")
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = T.SLIDE_W, T.SLIDE_H
        t_idx = c_idx = None

    with T.override(**overrides):
        pieces, fit_notes = expand(proposal)
        for num, msg in fit_notes:
            title = next((getattr(proposal.slides[num - 1], k) for k in ("title",) if getattr(proposal.slides[num - 1], k, None)), "")
            warnings.append(f"Slide {num}{f' ({title})' if title else ''}: {msg}.")
        for i, piece in enumerate(pieces, start=1):
            s = piece.slide
            o = piece.opts
            if template:
                layout = prs.slide_layouts[t_idx if s.type == "title" else c_idx]
            else:
                layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(layout)
            if s.type == "title":
                _title_slide(prs, slide, proposal.meta)
            elif s.type == "assumptions":
                a: Assumptions = s.assumptions
                _bullets_slide(slide, a.title, a.items, a.numbered, o.get("size", 15), o.get("start", 0))
            elif s.type == "bullets":
                _bullets_slide(slide, s.title, s.bullets, False, o.get("size", 15), o.get("start", 0))
            elif s.type == "gantt":
                _gantt_slide(slide, s.gantt, o.get("row_h"))
            elif s.type == "cost":
                _cost_slide(slide, s.cost, o.get("row_h"), o.get("lines"), o.get("final", True))
            elif s.type == "flow":
                _flow_slide(slide, s.flow)
            elif s.type == "solution":
                _solution_slide(slide, s.solution)
            elif s.type == "compare":
                _compare_slide(slide, s.compare)
            elif s.type == "scope":
                _scope_slide(slide, s.scope, o.get("size", 13))
            elif s.type == "risks":
                _risks_slide(slide, s.risks)
            elif s.type == "next_steps":
                _next_steps_slide(slide, s.next_steps, o.get("start", 0))
            if piece.part == 0 and s.notes.strip():
                slide.notes_slide.notes_text_frame.text = s.notes
            _describe_tables(slide, s)
            if template:
                _drop_empty_placeholders(slide)
                if s.type != "title":
                    _template_footer(slide, layout, proposal.meta, i)
            elif s.type != "title":
                _footer(slide, proposal.meta, i)
    if os.environ.get("STUDIO_EMBED_SOURCE", "1") != "0":  # lets a locked-down org opt out of the round-trip item
        interop.embed(prs, proposal)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
