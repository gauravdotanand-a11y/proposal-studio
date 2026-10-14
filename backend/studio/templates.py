"""Corporate template support: validate/convert uploads, inspect layouts, read theme colours."""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

MAX_UPLOAD = 25 * 1024 * 1024
MAX_UNCOMPRESSED = 300 * 1024 * 1024
POTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PPTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
NAME_RE = re.compile(r"^[A-Za-z0-9_-]{1,64}$")

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_AUX = {PH.DATE, PH.FOOTER, PH.SLIDE_NUMBER, PH.HEADER}


class TemplateError(ValueError):
    pass


def normalize(data: bytes) -> bytes:
    """Validate an uploaded .pptx/.potx and return plain .pptx bytes (python-pptx cannot open .potx)."""
    if len(data) > MAX_UPLOAD:
        raise TemplateError("file is larger than 25 MB")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise TemplateError("not a PowerPoint (.pptx/.potx) file") from e
    infos = zf.infolist()
    names = {i.filename for i in infos}
    if "ppt/presentation.xml" not in names or "[Content_Types].xml" not in names:
        raise TemplateError("not a PowerPoint presentation or template")
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise TemplateError("macro-enabled files are not supported; save as .potx or .pptx")
    if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
        raise TemplateError("file expands to an unreasonable size")
    ct = zf.read("[Content_Types].xml").decode("utf-8")
    fixed = ct.replace(POTX_CT, PPTX_CT)
    if fixed == ct:
        return data
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in infos:
            zo.writestr(i, fixed.encode("utf-8") if i.filename == "[Content_Types].xml" else zf.read(i.filename))
    return out.getvalue()


def _kinds(layout) -> set:
    return {p.placeholder_format.type for p in layout.placeholders}


def pick_layouts(prs) -> tuple[int, int]:
    """Best-guess (title_layout_index, content_layout_index)."""
    layouts = list(prs.slide_layouts)
    title_idx = 0
    for i, l in enumerate(layouts):
        k = _kinds(l)
        if PH.CENTER_TITLE in k or (PH.TITLE in k and PH.SUBTITLE in k):
            title_idx = i
            break
    content_idx = None
    for i, l in enumerate(layouts):
        if _kinds(l) - _AUX == {PH.TITLE}:
            content_idx = i
            break
    if content_idx is None:
        for i, l in enumerate(layouts):
            if PH.TITLE in _kinds(l) and i != title_idx:
                content_idx = i
                break
    return title_idx, content_idx if content_idx is not None else len(layouts) - 1


def _hex(el) -> str | None:
    if el is None:
        return None
    for child in el:
        tag = etree.QName(child).localname
        if tag == "srgbClr":
            return child.get("val")
        if tag == "sysClr":
            return child.get("lastClr")
    return None


def theme_colors(prs) -> dict[str, str]:
    master = prs.slide_masters[0]
    try:
        theme = master.part.part_related_by(RT.THEME)
    except KeyError:
        return {}
    root = etree.fromstring(theme.blob)
    scheme = root.find(f".//{{{_A}}}clrScheme")
    if scheme is None:
        return {}
    out = {}
    for name in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"):
        v = _hex(scheme.find(f"{{{_A}}}{name}"))
        if v:
            out[name] = v.upper()
    return out


def inspect(path: Path) -> dict:
    prs = Presentation(str(path))
    title_idx, content_idx = pick_layouts(prs)
    w_in, h_in = prs.slide_width / 914400, prs.slide_height / 914400
    warnings = []
    if abs(w_in / h_in - 16 / 9) > 0.03:
        warnings.append(f"Template is {w_in:.2f}x{h_in:.2f} in (not 16:9); diagrams and tables may be cramped.")
    return {
        "slide_w_in": round(w_in, 3), "slide_h_in": round(h_in, 3),
        "layouts": [{"index": i, "name": l.name,
                     "placeholders": sorted(p.placeholder_format.type.name for p in l.placeholders)}
                    for i, l in enumerate(prs.slide_layouts)],
        "title_layout": title_idx, "content_layout": content_idx,
        "colors": theme_colors(prs), "warnings": warnings,
    }


# --------------------------------------------------------------- palette
def _rgb(h: str) -> tuple[int, int, int]:
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _lum(h: str) -> float:
    r, g, b = _rgb(h)
    return (0.2126 * r + 0.7152 * g + 0.0722 * b) / 255


def _blend(h: str, white: float) -> str:
    return "".join(f"{round(c + (255 - c) * white):02X}" for c in _rgb(h))


def palette(colors: dict[str, str]) -> dict[str, str] | None:
    """Map template theme colours onto our roles; None when the theme has no usable accents."""
    a1, a2 = colors.get("accent1"), colors.get("accent2")
    if not a1:
        return None
    dk2 = colors.get("dk2")
    navy = dk2 if dk2 and _lum(dk2) < 0.45 else a1
    return {"NAVY": navy, "ORANGE": a2 or a1, "TEAL": _blend(a1, 0.8), "TEAL_DARK": _blend(a1, 0.45)}


def content_top_for(prs, layout_index: int) -> tuple[int, int]:
    """(left_margin, content_top) in EMU from the content layout's title placeholder."""
    layout = prs.slide_layouts[layout_index]
    for p in layout.placeholders:
        if p.placeholder_format.type in (PH.TITLE, PH.CENTER_TITLE) and p.height:
            left = min(max(int(p.left), Inches(0.3)), Inches(1.5))
            top = min(max(int(p.top + p.height + Inches(0.1)), Inches(1.0)), Inches(2.5))
            return left, top
    return int(Inches(0.6)), int(Inches(1.3))


# --------------------------------------------------------------- preview
_THEME_NAMES = {
    "BACKGROUND_1": "lt1", "TEXT_1": "dk1", "BACKGROUND_2": "lt2", "TEXT_2": "dk2",
    "ACCENT_1": "accent1", "ACCENT_2": "accent2", "ACCENT_3": "accent3",
    "ACCENT_4": "accent4", "ACCENT_5": "accent5", "ACCENT_6": "accent6",
    "LIGHT_1": "lt1", "DARK_1": "dk1", "LIGHT_2": "lt2", "DARK_2": "dk2",
}
MAX_PREVIEW_IMAGE = 400 * 1024


def _adjust(hexv: str, brightness: float) -> str:
    if not brightness:
        return hexv
    if brightness > 0:
        return _blend(hexv, min(brightness, 1.0))
    r, g, b = _rgb(hexv)
    k = 1 + max(brightness, -1.0)
    return f"{round(r * k):02X}{round(g * k):02X}{round(b * k):02X}"


def _fill_hex(fill, colors: dict[str, str]) -> str | None:
    """Solid fill colour as RRGGBB, or None (no fill / gradient / picture / unresolvable)."""
    try:
        from pptx.enum.dml import MSO_FILL, MSO_COLOR_TYPE
        if fill.type != MSO_FILL.SOLID:
            return None
        c = fill.fore_color
        if c.type == MSO_COLOR_TYPE.RGB:
            return str(c.rgb)
        if c.type == MSO_COLOR_TYPE.SCHEME:
            base = colors.get(_THEME_NAMES.get(c.theme_color.name, ""))
            return _adjust(base, c.brightness) if base else None
    except Exception:  # fill API raises for several unsupported element shapes
        return None
    return None


def _background(part, colors) -> str | None:
    try:
        if part._element.cSld.bg is None:
            return None
        return _fill_hex(part.background.fill, colors)
    except Exception:
        return None


def _shapes(container, colors) -> list[dict]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE as ST
    out = []
    for sh in container.shapes:
        try:
            x, y, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if None in (x, y, w, h) or not w or not h:
            continue
        item = {"x": int(x), "y": int(y), "w": int(w), "h": int(h), "kind": "shape", "label": "", "fill": None}
        if sh.is_placeholder:
            kind = sh.placeholder_format.type
            if kind in _AUX:
                item["kind"] = "aux"
            else:
                item["kind"] = "placeholder"
            item["label"] = kind.name.replace("_", " ").title() if kind else "Placeholder"
        elif sh.shape_type == ST.PICTURE:
            item["kind"] = "picture"
            try:
                img = sh.image
                if img.content_type in ("image/png", "image/jpeg") and len(img.blob) <= MAX_PREVIEW_IMAGE:
                    import base64
                    item["src"] = f"data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}"
            except Exception:
                pass
        elif sh.shape_type == ST.GROUP:
            item["kind"] = "group"
        else:
            try:
                item["fill"] = _fill_hex(sh.fill, colors)
            except Exception:
                pass
            if sh.has_text_frame and sh.text_frame.text.strip():
                item["label"] = sh.text_frame.text.strip()[:40]
        out.append(item)
    return out


def preview(path: Path) -> dict:
    """Schematic geometry of every layout (placeholders, background, master/layout decoration)."""
    prs = Presentation(str(path))
    colors = theme_colors(prs)
    master = prs.slide_masters[0]
    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        show_master = layout._element.get("showMasterSp", "1") not in ("0", "false")
        layouts.append({
            "index": i, "name": layout.name,
            "background": _background(layout, colors),
            "show_master_shapes": show_master,
            "shapes": _shapes(layout, colors),
        })
    return {
        "slide_w": int(prs.slide_width), "slide_h": int(prs.slide_height),
        "background": _background(master, colors) or colors.get("lt1", "FFFFFF"),
        "master_shapes": [s for s in _shapes(master, colors) if s["kind"] != "placeholder"],
        "layouts": layouts,
    }
