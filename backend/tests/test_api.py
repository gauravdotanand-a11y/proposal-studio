from fastapi.testclient import TestClient

from studio import api

client = TestClient(api.app)


def test_sample_roundtrip_and_render(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "DATA", tmp_path)
    p = client.get("/api/sample").json()
    assert client.post("/api/validate", json=p).json()["ok"]
    r = client.post("/api/render", json=p)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.put("/api/projects/demo", json=p).status_code == 200
    assert client.get("/api/projects").json() == ["demo"]
    assert client.get("/api/projects/demo").json()["meta"]["title"] == p["meta"]["title"]


def test_invalid_and_traversal(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "DATA", tmp_path)
    assert client.post("/api/validate", json={"meta": {}}).status_code == 422
    assert client.get("/api/projects/..%2Fetc").status_code in (400, 404)


def test_cost_and_weeks():
    r = client.post("/api/cost/summary", json={"years": ["Y0"], "lines": [{"head": "a", "values": [10]}], "benefits": [30]})
    assert r.json()["cba"] == 3.0
    r = client.post("/api/gantt/weeks", json={"start_date": "2026-11-02", "activities": [{"name": "x", "start_week": 1, "end_week": 2}]})
    assert r.json()[1]["label"] == "9 Nov - 13 Nov"


def test_flow_layout_matches_export(tmp_path):
    from pptx import Presentation
    from studio.models import Proposal
    p = client.get("/api/sample").json()
    flow = next(s["flow"] for s in p["slides"] if s["type"] == "flow")
    lay = client.post("/api/flow/layout", json=flow).json()
    assert len(lay["nodes"]) == len(flow["nodes"]) and len(lay["edges"]) == len(flow["edges"])
    out = tmp_path / "o.pptx"
    from studio.render import render
    render(Proposal.model_validate(p), out)
    slide = Presentation(out).slides[1]
    boxes = {(s.left, s.top, s.width, s.height) for s in slide.shapes if s.has_text_frame and s.shape_type.name == "AUTO_SHAPE"}
    for n in lay["nodes"]:
        assert (n["x"], n["y"], n["w"], n["h"]) in boxes
    bad = client.post("/api/flow/layout", json={"title": "t", "nodes": [], "edges": [{"src": "a", "dst": "b", "label": ""}]})
    assert bad.status_code == 422


def test_validate_reports_fit_notes():
    from fastapi.testclient import TestClient
    from studio.api import app
    long = "Extremely long label text that keeps going to test wrapping inside boxes " * 2
    body = {"meta": {"title": "T"}, "slides": [{"type": "assumptions", "assumptions": {
        "title": "A", "items": [long + str(i) for i in range(14)], "numbered": True}}]}
    r = TestClient(app).post("/api/validate", json=body).json()
    assert r["ok"] and r["notes"] and r["notes"][0]["slide"] == 1
