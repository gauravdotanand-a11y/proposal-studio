from datetime import date
from pathlib import Path

import pytest
import yaml
from pptx import Presentation

from studio.models import Cost, CostLine, Flow, Gantt, Activity, Proposal
from studio.render import render

SAMPLE = Path(__file__).parents[2] / "examples" / "sample.yaml"


def test_cba_and_totals():
    c = Cost(lines=[CostLine(head="a", values=[10, 5]), CostLine(head="b", values=[0, 5])],
             years=["Y0", "Y1"], benefits=[0, 40])
    assert c.year_totals == [10, 10]
    assert c.total_cost == 20 and c.cba == 2.0


def test_cost_length_mismatch():
    with pytest.raises(ValueError):
        Cost(lines=[CostLine(head="a", values=[1])], benefits=[1, 2, 3, 4])


def test_week_range_mon_fri():
    g = Gantt(start_date=date(2026, 11, 4), activities=[Activity(name="x", start_week=1, end_week=2)])
    s, e = g.week_range(2)
    assert (s, e) == (date(2026, 11, 9), date(2026, 11, 13))
    assert g.week_count == 2


def test_flow_unknown_edge():
    with pytest.raises(ValueError):
        Flow(title="t", nodes=[], edges=[{"src": "a", "dst": "b"}])


def test_render_sample(tmp_path):
    p = Proposal.model_validate(yaml.safe_load(SAMPLE.read_text()))
    out = render(p, tmp_path / "o.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == 12
    flow_slide = prs.slides[1]
    edges = [s for s in flow_slide.shapes if s.name.startswith("edge:")]
    assert len(edges) == 8


def test_solution_slide_renders_and_validates():
    import pytest as _pt
    from studio.models import Solution
    p = Proposal.model_validate(yaml.safe_load(SAMPLE.read_text()))
    import tempfile
    out = render(p, Path(tempfile.mkdtemp()) / "s.pptx")
    prs = Presentation(out)
    sol = next(s for s in prs.slides if any(sh.name.startswith("stage:") for sh in s.shapes))
    stages = [sh for sh in sol.shapes if sh.name.startswith("stage:")]
    assert len(stages) == 6 and len([sh for sh in sol.shapes if sh.name == "stage-arrow"]) == 5
    assert all(sh.left + sh.width <= prs.slide_width for sh in stages)
    text = " ".join(sh.text_frame.text for sh in sol.shapes if sh.has_text_frame)
    assert "High Level End to End Flow" in text and "1. Fetch:" in text
    with _pt.raises(ValueError):
        Solution(stages=[])
    with _pt.raises(ValueError):
        Solution(stages=[{"title": str(i)} for i in range(8)])


def test_swimlane_slide_has_lanes_badges_and_note():
    p = Proposal.model_validate(yaml.safe_load(SAMPLE.read_text()))
    import tempfile
    prs = Presentation(render(p, Path(tempfile.mkdtemp()) / "l.pptx"))
    s = next(s for s in prs.slides if any(sh.name.startswith("lane:") for sh in s.shapes))
    names = [sh.name for sh in s.shapes]
    assert "lane:RPA Bot" in names and "lane:Agent" in names
    assert any(n.startswith("badge:") for n in names) and any(n.startswith("note:") for n in names)


def test_week_label_skips_holidays():
    from datetime import date
    from studio.models import Gantt
    g = Gantt(start_date=date(2026, 12, 21), activities=[{"name": "a", "start_week": 1, "end_week": 1}],
              holidays=[date(2026, 12, 21), date(2026, 12, 25)])
    assert g.week_range(1) == (date(2026, 12, 22), date(2026, 12, 24))
