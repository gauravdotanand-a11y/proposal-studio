import tempfile
from datetime import date
from pathlib import Path

import pytest
import yaml
from pptx import Presentation

from studio.models import Activity, Compare, Cost, CostLine, Driver, Gantt, Proposal, Risks
from studio.render import render

SAMPLE = Path(__file__).parents[2] / "examples" / "sample.yaml"


def _deck():
    p = Proposal.model_validate(yaml.safe_load(SAMPLE.read_text()))
    return p, Presentation(render(p, Path(tempfile.mkdtemp()) / "d.pptx"))


def _gantt(*acts, start=date(2026, 11, 2)):
    return Gantt(start_date=start, activities=[Activity(name=n, start_week=s, end_week=e) for n, s, e in acts])


def test_driver_math_and_ramp():
    d = Driver(name="x", annual_volume=60000, minutes_saved=3, hourly_rate=20, success_rate=0.5, ramp=[0.1, 1])
    # 60000 * 3 / 60 * 20 * 0.5 = 30000 per full year
    assert d.values(2) == [3000.0, 30000.0]
    assert Driver(name="y", annual_volume=1, minutes_saved=60, hourly_rate=10).values(3) == [10.0, 10.0, 10.0]


def test_cost_with_drivers_and_manual_benefits():
    c = Cost(years=["Y0", "Y1"], lines=[CostLine(head="dev", values=[100, 0], kind="one_time"),
                                        CostLine(head="run", values=[10, 10], kind="recurring")],
             benefits=[5, 5], drivers=[Driver(name="d", annual_volume=60, minutes_saved=60, hourly_rate=1, ramp=[0, 1])])
    assert c.benefit_values == [5.0, 65.0]
    assert c.total_cost == 120 and c.one_time_total == 100 and c.recurring_total == 20
    assert c.cba == pytest.approx(70 / 120)


def test_cost_validation():
    with pytest.raises(ValueError):
        Cost(lines=[CostLine(head="a", values=[1, 1, 1, 1])])  # no benefits and no drivers
    with pytest.raises(ValueError):
        Cost(lines=[CostLine(head="a", values=[1, 1, 1, 1])],
             drivers=[Driver(name="d", annual_volume=1, minutes_saved=1, hourly_rate=1, ramp=[1])])


def test_compare_changed_and_delta():
    c = Compare(previous=_gantt(("A", 1, 1), ("B", 2, 4)), updated=_gantt(("A", 1, 1), ("B", 2, 5), ("C", 6, 6)))
    assert c.changed == ["B", "C"] and c.delta_weeks == 2


def test_new_slides_render_with_expected_shapes():
    p, prs = _deck()
    slides = list(prs.slides)
    types = [s.type for s in p.slides]
    cmp_slide = slides[types.index("compare")]
    tables = [sh for sh in cmp_slide.shapes if sh.has_table]
    assert len(tables) == 2 and {t.name for t in tables} == {"gantt:Previous version", "gantt:Updated version (4 Aug)"}
    text = " ".join(sh.text_frame.text for sh in cmp_slide.shapes if sh.has_text_frame)
    assert "+1 week(s)" in text and "4 activities changed" in text

    risks = next(sh for sh in slides[types.index("risks")].shapes if sh.name == "risks-table").table
    assert str(risks.cell(1, 1).fill.fore_color.rgb) == "FF4F1F"  # High -> orange
    assert risks.cell(1, 1).text_frame.text == "High"

    scope = slides[types.index("scope")]
    assert {"scope-head:0", "scope-head:1"} <= {sh.name for sh in scope.shapes}
    nxt = next(sh for sh in slides[types.index("next_steps")].shapes if sh.name == "next-steps-table").table
    assert len(nxt.rows) == 4 and nxt.cell(1, 1).text_frame.text.startswith("Confirm scope")


def test_gantt_sections_become_bands():
    p, prs = _deck()
    s = [x for x, t in zip(prs.slides, [s.type for s in p.slides]) if t == "gantt"][1]
    tbl = next(sh for sh in s.shapes if sh.has_table).table
    first_col = [tbl.cell(r, 0).text_frame.text for r in range(len(tbl.rows))]
    assert "New error codes" in first_col and "Reporting and INC monitoring" in first_col
    r = first_col.index("New error codes")
    assert tbl.cell(r, 0).is_merge_origin


def test_cost_slide_shows_driver_basis_and_totals():
    p, prs = _deck()
    s = list(prs.slides)[[x.type for x in p.slides].index("cost")]
    text = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    assert "Invoice data entry: 120,000 items/yr x 4 min" in text and "One-time: $75,000" in text


def test_nothing_runs_off_the_slide():
    p, prs = _deck()
    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            h = sh.height
            if sh.has_table:
                h = sum(r.height for r in sh.table.rows)
            assert sh.left >= 0 and sh.left + sh.width <= prs.slide_width + 10, f"slide {i}: {sh.name} off right"
            assert sh.top + h <= prs.slide_height + 10, f"slide {i}: {sh.name} off bottom ({sh.top + h} > {prs.slide_height})"


def test_risks_requires_valid_impact():
    with pytest.raises(ValueError):
        Risks.model_validate({"items": [{"risk": "x", "impact": "Huge"}]})
