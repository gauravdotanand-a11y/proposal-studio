import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

from studio import api, templates

client = TestClient(api.app)


def _make_template(widescreen=True) -> bytes:
    prs = Presentation()  # python-pptx default template: has "Title Slide", "Title Only", ...
    if widescreen:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # a leftover sample slide that must not leak into exports
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "TEMPLATE SAMPLE"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _as_potx(data: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zi, zipfile.ZipFile(out, "w") as zo:
        for i in zi.infolist():
            b = zi.read(i.filename)
            if i.filename == "[Content_Types].xml":
                b = b.replace(templates.PPTX_CT.encode(), templates.POTX_CT.encode())
            zo.writestr(i, b)
    return out.getvalue()


@pytest.fixture()
def dirs(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "TEMPLATES", tmp_path / "t")
    monkeypatch.setattr(api, "DATA", tmp_path / "p")
    return tmp_path


def test_potx_upload_inspect_and_render(dirs):
    r = client.put("/api/templates/corp", content=_as_potx(_make_template()))
    assert r.status_code == 200, r.text
    info = r.json()
    names = [l["name"] for l in info["layouts"]]
    assert names[info["title_layout"]] == "Title Slide" and names[info["content_layout"]] == "Title Only"
    assert "accent1" in info["colors"]
    assert client.get("/api/templates").json() == ["corp"]

    deck = client.get("/api/sample").json()
    out = client.post("/api/render?template=corp", json=deck)
    assert out.status_code == 200 and out.content[:2] == b"PK"
    prs = Presentation(io.BytesIO(out.content))
    assert len(prs.slides) == len(deck["slides"])  # template's sample slide removed
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert prs.slides[0].shapes.title.text_frame.text == deck["meta"]["title"]
    assert all(s.slide_layout.name == "Title Only" for s in list(prs.slides)[1:])
    assert prs.slides[1].shapes.title.text_frame.text == "HL Flow of Automation"
    assert not any("TEMPLATE SAMPLE" in sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


def test_theme_colors_applied_and_optional(dirs):
    client.put("/api/templates/corp", content=_make_template())
    deck = client.get("/api/sample").json()
    gantt = lambda r: next(s for s in Presentation(io.BytesIO(r.content)).slides if any(sh.has_table for sh in s.shapes))
    on = client.post("/api/render?template=corp", json=deck)
    off = client.post("/api/render?template=corp&use_theme_colors=false", json=deck)
    def header_fill(r):
        tbl = next(sh for sh in gantt(r).shapes if sh.has_table).table
        return str(tbl.cell(0, 1).fill.fore_color.rgb)
    assert header_fill(off) == "1B2A7A"
    assert header_fill(on) != "1B2A7A"


def test_43_template_warns(dirs):
    client.put("/api/templates/old", content=_make_template(widescreen=False))
    out = client.post("/api/render?template=old", json=client.get("/api/sample").json())
    assert out.status_code == 200 and "not 16:9" in out.headers["x-render-warnings"]


def test_rejects_bad_uploads(dirs):
    assert client.put("/api/templates/x", content=b"not a zip").status_code == 400
    bad = io.BytesIO()
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/presentation.xml", "<p/>")
        z.writestr("ppt/vbaProject.bin", "x")
    assert "macro" in client.put("/api/templates/x", content=bad.getvalue()).json()["detail"]
    assert client.put("/api/templates/a.b", content=b"x").status_code == 400
    assert client.put("/api/templates/..%2Fevil", content=b"x").status_code in (400, 404, 405)
    assert not list((dirs / "t").glob("*evil*")) if (dirs / "t").exists() else True
    assert client.post("/api/render?template=missing", json=client.get("/api/sample").json()).status_code == 404
    client.put("/api/templates/corp", content=_make_template())
    assert client.post("/api/render?template=corp&content_layout=99", json=client.get("/api/sample").json()).status_code == 400


def test_preview_geometry(dirs):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # python-pptx cannot add shapes to a master: build on a scratch slide, then move the element
    scratch = prs.slides.add_slide(prs.slide_layouts[6])
    bar = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    prs.slide_masters[0].shapes._spTree.append(bar._element)
    buf = io.BytesIO()
    prs.save(buf)
    assert client.put("/api/templates/p", content=buf.getvalue()).status_code == 200
    r = client.get("/api/templates/p/preview")
    assert r.status_code == 200
    d = r.json()
    assert d["slide_w"] == int(Inches(13.333)) and len(d["layouts"]) == 11
    assert any(s["fill"] == "123456" for s in d["master_shapes"])
    title_only = next(l for l in d["layouts"] if l["name"] == "Title Only")
    phs = [s for s in title_only["shapes"] if s["kind"] == "placeholder"]
    assert any(s["label"] == "Title" and s["w"] > 0 and s["h"] > 0 for s in phs)
    assert client.get("/api/templates/nope/preview").status_code == 404


def test_presets_render_with_their_colors_and_font(tmp_path):
    import yaml
    from pptx import Presentation
    from studio import theme
    from studio.models import Proposal
    from studio.render import render
    p = Proposal.model_validate(yaml.safe_load(open("../examples/sample.yaml")))
    seen = {}
    for name in theme.PRESETS:
        out = tmp_path / f"{name}.pptx"
        render(p, out, preset=name)
        prs = Presentation(str(out))
        fills = set()
        for sh in prs.slides[1].shapes:
            try:
                fills.add(str(sh.fill.fore_color.rgb))
            except Exception:
                pass
        seen[name] = fills
    assert theme.PRESETS["forest"]["colors"]["TEAL"] in seen["forest"]
    assert theme.PRESETS["corporate"]["colors"]["TEAL"] not in seen["forest"]
    assert theme.FONT == "Arial"  # override restored after rendering


def test_unknown_preset_rejected(tmp_path):
    import pytest, yaml
    from studio.models import Proposal
    from studio.render import render
    p = Proposal.model_validate(yaml.safe_load(open("../examples/sample.yaml")))
    with pytest.raises(ValueError):
        render(p, tmp_path / "x.pptx", preset="nope")


def test_preset_recolours_body_text(tmp_path):
    import yaml
    from pptx import Presentation
    from studio import theme
    from studio.models import Proposal
    from studio.render import render
    p = Proposal.model_validate(yaml.safe_load(open("../examples/sample.yaml")))
    out = tmp_path / "m.pptx"
    render(p, out, preset="mono")
    colors = set()
    for sl in Presentation(str(out)).slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        try:
                            colors.add(str(r.font.color.rgb))
                        except Exception:
                            pass
    assert theme.PRESETS["mono"]["colors"]["GREY"] in colors
    assert theme.PRESETS["corporate"]["colors"]["GREY"] not in colors
